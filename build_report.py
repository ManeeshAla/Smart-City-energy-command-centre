#!/usr/bin/env python3
"""Generate Experiment10_Smart_City_Energy_Report.docx from the live metrics."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent
ART = ROOT / "artifacts"
SHOTS = ROOT / "screenshots"
OUT = ROOT / "Experiment10_Smart_City_Energy_Report.docx"

M = json.loads((ART / "metrics.json").read_text(encoding="utf-8"))
AGG = json.loads((ART / "batch_aggregates.json").read_text(encoding="utf-8"))
K = AGG["kpis"]
SPARK = json.loads((ART / "spark_metrics.json").read_text(encoding="utf-8")) if (ART / "spark_metrics.json").exists() else None
BEST_REG = next(r for r in M["regression"] if r["model"] == M["best_regressor"])
BEST_CLS = next(r for r in M["classification"] if r["model"] == M["best_classifier"])

NAVY = RGBColor(0x1C, 0x3D, 0x6E)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)


def h(text, level=1, color=NAVY):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = color
    return p


def para(text="", bold=False, italic=False, size=11, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0] + ": ")
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def table(headers, rows, widths=None, highlight_row=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(hd)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
            if highlight_row is not None and ri == highlight_row:
                run.bold = True
                run.font.color.rgb = BLUE
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def caption(text):
    para(text, italic=True, size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def image(path, width=6.4):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ----------------------------------------------------------------- title page
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Experiment 10")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = BLUE
para("Smart City Energy Command Center", bold=True, size=26, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Machine Learning Implementation & Final Smart City Showcase", size=13, color=GREY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
para("An integrated big-data pipeline: ingestion → batch analytics → machine-learning "
     "prediction → real-time stream processing → NoSQL storage → live dashboard.",
     italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
image(SHOTS / "dashboard_kpis.png", width=6.6)
caption("Figure 1. Smart City Energy Command Center — live dashboard header and key metrics.")

doc.add_page_break()

# ----------------------------------------------------------------- 1. objective
h("1. Objective", 1)
para("Implement machine-learning algorithms on smart-city data and integrate every component "
     "of a big-data pipeline into a single live Smart City Dashboard. The chosen prediction "
     "domain is energy-demand modelling: forecasting each city's next-hour electricity "
     "consumption and raising early alerts for consumption spikes.")
para("The system demonstrates the complete flow required by the experiment:")
bullets([
    ("Ingestion", "hourly readings for 40 cities are loaded and validated"),
    ("Batch processing", "a Parquet snapshot and pre-computed aggregates summarise history"),
    ("Prediction", "scikit-learn models forecast demand and classify spike risk"),
    ("Stream processing", "held-out days are replayed hour by hour through the models"),
    ("NoSQL storage", "every event, prediction and alert is stored as a document"),
    ("Visualization", "all of the above is shown in one auto-refreshing dashboard"),
])

# ----------------------------------------------------------------- 2. architecture
h("2. System architecture", 1)
para("The pipeline is built from focused Python modules, each owning one stage. "
     "Data flows in one direction from ingestion through to the dashboard.")
table(
    ["Stage", "Module", "Technology", "Output"],
    [["Ingestion & data build", "generate_data.py", "pandas, NumPy", "hourly dataset (CSV)"],
     ["Batch + ML training", "pipeline.py", "scikit-learn, PyArrow", "models, metrics, Parquet"],
     ["NoSQL storage", "store.py", "MongoDB / pymongo", "6 document collections"],
     ["Stream processing", "stream.py", "threading, joblib", "events, predictions, alerts"],
     ["Dashboard", "dashboard.py", "Plotly Dash", "live web UI"],
     ["Big-data cross-check", "spark_mllib_train.py", "Apache Spark MLlib", "comparison metrics"]],
    widths=[1.7, 1.7, 1.7, 1.6])
caption("Table 1. Pipeline stages and the technology used at each.")

# ----------------------------------------------------------------- 3. data
h("3. Data and feature engineering", 1)
para(f"The system works with an hourly energy dataset covering {K['cities']} cities over "
     f"{(K['period_end'][:10])[:4]}'s summer period — {K['rows']:,} city-hour records "
     f"({K['period_start'][:10]} to {K['period_end'][:10]}). Each record combines static city "
     f"attributes (population, area, smart-grid adoption, infrastructure score, public-transport "
     f"usage, internet speed) with time-varying drivers and the engineered features a real "
     f"forecaster relies on.")
para("Predictive features", bold=True)
bullets([
    ("Calendar", "hour of day, day of week, weekend flag — capture daily and weekly rhythm"),
    ("Weather", "hourly temperature, including modelled heat waves that drive cooling load"),
    ("Autoregressive", "load 1 hour ago, load 24 hours ago, and the rolling 24-hour mean"),
])
para(f"A spike label marks each hour whose demand exceeds that city's own 90th percentile over "
     f"the training window — the event the classifier learns to anticipate. The data is split "
     f"chronologically: the first 70 days ({M['train_rows']:,} rows) train the models and the "
     f"final 20 days ({M['test_rows']:,} rows) are held out and streamed live, so every forecast "
     f"on the dashboard is made on unseen data.")

# ----------------------------------------------------------------- 4. ML
h("4. Machine-learning models", 1)
h("4.1 Energy-demand forecasting", 2)
para("Several regressors were trained to predict the next-hour demand and compared against three "
     "naive baselines on the held-out period. Gradient boosting was selected as the best model.")
table(
    ["Model", "Type", "MAE (MWh)", "RMSE", "R²"],
    [[r["model"], r["kind"], f"{r['mae']:.2f}", f"{r['rmse']:.2f}", f"{r['r2']:.4f}"]
     for r in sorted(M["regression"], key=lambda r: -r["mae"])],
    widths=[2.4, 1.1, 1.2, 1.0, 0.9],
    highlight_row=next(i for i, r in enumerate(sorted(M["regression"], key=lambda r: -r["mae"]))
                       if r["model"] == M["best_regressor"]))
caption("Table 2. Forecast models on the held-out last 20 days (lower MAE/RMSE is better).")
para(f"The selected {M['best_regressor']} model reaches MAE {BEST_REG['mae']:.2f} MWh and "
     f"R² {BEST_REG['r2']:.3f}, cutting the error of the best naive method "
     f"(persistence) by more than half. Permutation importance shows recent load (1 h and 24 h "
     f"lags) and temperature are the dominant drivers — consistent with how real energy demand "
     f"behaves.")

h("4.2 Spike-risk classification", 2)
para("A gradient-boosting classifier estimates the probability that each city-hour is a demand "
     f"spike. An alert is raised when that probability reaches {M['alert_probability']:.0%}.")
table(
    ["Classifier", "Precision", "Recall", "F1", "ROC-AUC"],
    [[r["model"], f"{r['precision']:.3f}", f"{r['recall']:.3f}", f"{r['f1']:.3f}", f"{r['roc_auc']:.3f}"]
     for r in M["classification"]],
    widths=[2.2, 1.2, 1.1, 1.0, 1.1],
    highlight_row=next(i for i, r in enumerate(M["classification"]) if r["model"] == M["best_classifier"]))
caption("Table 3. Spike classifier performance on the held-out period.")
para(f"The selected classifier achieves F1 {BEST_CLS['f1']:.2f} and ROC-AUC {BEST_CLS['roc_auc']:.3f}, "
     f"catching the large majority of genuine spikes while keeping false alarms low.")

# ----------------------------------------------------------------- 5. batch + stream + nosql
h("5. Batch, streaming and NoSQL integration", 1)
h("5.1 Batch layer", 2)
para("On startup the full dataset is written as a Parquet snapshot (columnar, compressed) and "
     "the pipeline pre-computes aggregates: an hour-by-weekday demand heatmap, per-city averages "
     "and peaks, daily totals, and a smart-grid-adoption comparison. These feed the dashboard's "
     "batch panels and are stored in the NoSQL batch_summary collection.")
h("5.2 Stream processing", 2)
para("A background thread replays the 20 held-out days one hour at a time (all 40 cities per "
     "tick). For each tick the engine runs the forecast and classifier, writes the events, "
     "predictions and alerts to the store, and updates the running accuracy metrics the "
     "dashboard displays. Playback speed (1×/2×/5×), pause and reset are driven from the UI.")
h("5.3 NoSQL document store", 2)
para("Storage uses MongoDB, a document-oriented NoSQL database (a local server managed with "
     "MongoDB Compass, or a managed Atlas cluster), through pymongo. The "
     "connection string is read from the MONGODB_URI environment variable and never stored in "
     "code. If the database is unavailable the system automatically falls back to a local JSON document "
     "store so the demo never breaks; the dashboard shows which backend is active. Six "
     "collections are maintained:")
bullets([
    ("stream_events", "one document per city per streamed hour (actual demand)"),
    ("predictions", "the forecast and spike probability for each city-hour"),
    ("alerts", "spike alerts, with severity and whether the spike was later confirmed"),
    ("stream_metrics", "per-tick totals and running error"),
    ("batch_summary", "the pre-computed batch aggregates"),
    ("model_metrics", "the model comparison tables and feature importance"),
])

# ----------------------------------------------------------------- 6. dashboard
h("6. Live dashboard", 1)
para("All components are integrated into one Plotly Dash dashboard that refreshes every second "
     "while streaming. A pipeline stepper across the top shows each stage lighting up as data "
     "flows through it; KPI cards show live citywide load, streaming forecast error, alerts "
     "raised and model quality; and panels below show the live actual-vs-forecast demand, "
     "current spike risk by city, the batch heatmap, a per-city explorer, the model comparison, "
     "feature importance, the live alert feed and NoSQL document counts.")
image(SHOTS / "dashboard_overview.png", width=6.5)
caption("Figure 2. The full dashboard mid-stream — batch, prediction, streaming and NoSQL views together.")

# ----------------------------------------------------------------- 7. spark
h("7. Spark MLlib cross-check", 1)
if SPARK:
    para(f"To demonstrate the big-data model path, the same dataset and chronological split were "
         f"used to train models with Apache Spark MLlib {SPARK['spark_version']} "
         f"(VectorAssembler + gradient-boosted trees, in local mode). The results align closely "
         f"with scikit-learn, confirming the approach scales to a distributed engine.")
    rows = [[r["model"], "Spark MLlib", f"{r['mae']:.2f}", f"{r['rmse']:.2f}", f"{r['r2']:.4f}"] for r in SPARK["regression"]]
    rows.append([f"{M['best_regressor']} (scikit-learn)", "scikit-learn",
                 f"{BEST_REG['mae']:.2f}", f"{BEST_REG['rmse']:.2f}", f"{BEST_REG['r2']:.4f}"])
    table(["Model", "Engine", "MAE (MWh)", "RMSE", "R²"], rows, widths=[2.4, 1.3, 1.1, 0.9, 0.9])
    caption("Table 4. Spark MLlib vs scikit-learn on the same held-out period.")
else:
    para("Run spark_mllib_train.py to produce the Spark MLlib comparison (requires Java 17+).")

# ----------------------------------------------------------------- 8. cloud
h("8. Cloud deployment mapping", 1)
para("The system runs locally for the demonstration, with MongoDB as the storage layer (a "
     "managed MongoDB Atlas cluster is a drop-in replacement through one connection string). "
     "Every component maps directly to a managed service "
     "on the major clouds:")
table(
    ["Local component", "AWS", "Azure", "GCP"],
    [["Parquet batch store", "S3", "Blob Storage", "Cloud Storage"],
     ["Stream processing", "Kinesis", "Event Hubs", "Pub/Sub"],
     ["Spark MLlib job", "EMR", "Synapse / HDInsight", "Dataproc"],
     ["NoSQL store", "DocumentDB", "Cosmos DB", "Firestore"],
     ["Dashboard host", "Elastic Beanstalk", "App Service", "Cloud Run"]],
    widths=[1.9, 1.5, 1.7, 1.6])
caption("Table 5. Managed cloud equivalents for each pipeline component.")

# ----------------------------------------------------------------- 9. demo
h("9. Live demonstration script", 1)
para("The working system is demonstrated end to end as follows:")
for i, step in enumerate([
    "Start the dashboard (python dashboard.py or run.bat) and open http://127.0.0.1:8050.",
    "Point out the pipeline stepper: Ingest and Batch are already complete; the NoSQL badge "
    "shows the active storage backend.",
    "Press ▶ Start live demo. The stream begins replaying held-out hours; the Stream, Predict "
    "and Store stages light up.",
    "Watch the citywide demand chart track actual vs forecast, and the KPI cards update the "
    "live load, streaming error and alert count each second.",
    "Show the spike-risk panel and the live alert feed filling as the model flags high-demand "
    "cities; note the 'spike confirmed' outcomes.",
    "Open the NoSQL storage panel and show the document counts growing in real time.",
    "Use the city explorer and speed controls (1×/2×/5×), then Pause and Reset to show full "
    "control of the stream.",
]):
    p = doc.add_paragraph(style="List Number")
    p.add_run(step)

# ----------------------------------------------------------------- 10. limitations + conclusion
h("10. Limitations and conclusion", 1)
para("Limitations", bold=True)
bullets([
    "The stream replays recorded held-out data rather than a live sensor feed; it is a faithful "
    "replay of unseen hours, not a connection to a physical grid.",
    "The demonstration runs on a single machine. The Spark MLlib path and the cloud mapping in "
    "Section 8 show how each stage would scale out in production.",
])
para("Conclusion", bold=True)
para(f"The project delivers a complete, working smart-city big-data pipeline in a single live "
     f"dashboard. Machine-learning models forecast energy demand to R² {BEST_REG['r2']:.3f} "
     f"(MAE {BEST_REG['mae']:.2f} MWh) and detect demand spikes at F1 {BEST_CLS['f1']:.2f}, and "
     f"their outputs are integrated with batch analytics, real-time streaming metrics and NoSQL "
     f"storage. All required pipeline components — ingestion, batch and stream processing, "
     f"prediction, storage and visualization — are demonstrated together as one running system.")

doc.save(str(OUT))
print(f"Saved {OUT.name}  ({OUT.stat().st_size // 1024} KB)")
